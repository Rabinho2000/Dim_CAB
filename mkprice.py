"""Preços de cabo: folhas JMM e Indusmelec Total, só compras de 2026, a mais recente de cada cabo (entre os dois fornecedores)."""
import openpyxl,datetime,re,json
wb=openpyxl.load_workbook('jmm/base.xlsx',read_only=True,data_only=True)
rows=[]
for r in list(wb['JMM'].iter_rows(values_only=True))[1:]:
    if isinstance(r[4],datetime.datetime): rows.append(('JMM',str(r[1]),str(r[2] or ''),float(r[3]),r[4]))
for r in list(wb['Indusmelec Total'].iter_rows(values_only=True))[1:]:
    if isinstance(r[9],datetime.datetime) and r[7] is not None: rows.append(('Indusmelec',str(r[1]),str(r[2] or ''),float(r[7]),r[9]))
rows=[x for x in rows if x[4].year>=2026]
desc={}
for f,ref,d,p,dt in rows:
    if 'CABO' in d.upper() and len(d.strip())>6: desc[ref]=d
def parse(ref,d):
    d=(desc.get(ref,d) if len(d.strip())<=6 else d).upper().replace(',','.')
    # terminais: bimetálicos Cu/Al (C1AU70, TBF 95, CAUBT95, "BIMET 95") e de cravar em cobre ("TERMINAL 25-10", "L 150/12MS");
    # fora: mangas termorretráteis, ponteiras isoladas e ferramentas
    if re.search(r'TERM|BIMET',d) and not re.search(r'TERMOR|AISL|PONTEIR|MATRIZ',d):
        m=re.search(r'\bC\dAU(\d+)\b|BIMET\w*\s+(\d+)|\bTBF\s*(\d+)|CAUBT(\d+)',d)
        if m: return ('Al','term',float(next(g for g in m.groups() if g)))
        m=re.search(r'\b(\d{2,3})[-/]\d{1,2}(?:MS)?\b',d)
        if m and 'TERM' in d: return ('Cu','term',float(m.group(1)))
        return None
    if re.search(r'H07V-K.*\bV/A|FV 6 .*VD/AM',d):
        m=re.search(r'(\d+(?:\.\d+)?) (?:V/A|LIVRE)',d);return ('Cu','pe',float(m.group(1))) if m else None
    m=re.search(r'\bLXV(?:/XZ1)? 1X(\d+)',d)
    if m: return ('Al','mono',float(m.group(1)))
    m=re.search(r'\b(?:RZ1-K|XV|RV-?K|XZ1)\s+1X(\d+(?:\.\d+)?)',d)
    if m: return ('Cu','mono',float(m.group(1)))
    m=re.search(r'\b(?:RZ1-?K|XV|RV-?K|RVK)\s+([345])G(\d+(?:\.\d+)?)',d)
    if m and m.group(1)!='4': return ('Cu','multi'+m.group(1),float(m.group(2)))
    return None
best={}
for f,ref,d,p,dt in rows:
    k=parse(ref,d)
    if not k: continue
    des=(desc.get(ref,d)).strip()
    if k[1]=='term': des=next((l.strip() for l in des.splitlines() if 'TERM' in l.upper() or 'BIMET' in l.upper()),des)  # guias de transporte com várias linhas
    if k not in best or dt>best[k][1]: best[k]=(p,dt,ref,des,f)
out=[[k[0],k[1],k[2],round(v[0],4),v[1].strftime('%Y-%m-%d'),v[2],v[3],v[4]] for k,v in sorted(best.items())]
json.dump(out,open('jmm/prices.json','w',encoding='utf-8'),ensure_ascii=False)
for o in out: print(o)
